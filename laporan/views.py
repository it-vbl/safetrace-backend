import io
import json
from decimal import Decimal

import openpyxl
from django.core.files.base import ContentFile
from django.http import FileResponse, Http404, HttpResponse
from django.db import transaction
from django.db.models import Q
from django.shortcuts import get_object_or_404
from openpyxl.styles import Font
from rest_framework import generics, status
from rest_framework.response import Response
from rest_framework_simplejwt.authentication import JWTAuthentication
from rest_framework.permissions import IsAuthenticated
from kebun.models import Kebun
from petani.models import Petani, Pekerja, Diklat
from penjualan.models import Angkutan, KelompokPenyetor
from .models import LaporanStatistik
from .serializers import (
    LaporanPetaniSerializer,
    LaporanStatistikCreateInputSerializer,
    LaporanStatistikSerializer,
)
from utils.reports import generate_laporan_statistik_kelompok_excel
from utils.decorators import role_required

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class LaporanStatistikListView(generics.ListAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    serializer_class = LaporanStatistikSerializer
    queryset = LaporanStatistik.objects.all()

    def get_queryset(self):
        queryset = super().get_queryset()
        search = (self.request.query_params.get('search') or '').strip()

        if not search:
            return queryset

        return queryset.filter(Q(judul__icontains=search) | Q(kebutuhan__icontains=search))


class LaporanStatistikDetailView(generics.RetrieveAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    serializer_class = LaporanStatistikSerializer
    queryset = LaporanStatistik.objects.all()


class LaporanStatistikDownloadView(generics.GenericAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    queryset = LaporanStatistik.objects.all()

    def get(self, request, *args, **kwargs):
        laporan = self.get_object()
        if not laporan.file_excel:
            raise Http404('File excel tidak ditemukan untuk laporan ini.')

        filename = laporan.file_excel.name.split('/')[-1]
        return FileResponse(
            laporan.file_excel.open('rb'),
            as_attachment=True,
            filename=filename,
            content_type=XLSX_CONTENT_TYPE,
        )


class LaporanStatistikDeleteView(generics.DestroyAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    queryset = LaporanStatistik.objects.all()

    def perform_destroy(self, instance):
        if instance.file_excel:
            storage = instance.file_excel.storage
            file_name = instance.file_excel.name
            instance.delete()
            if storage.exists(file_name):
                storage.delete(file_name)
        else:
            instance.delete()


class LaporanStatistikCreateView(generics.GenericAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    serializer_class = LaporanStatistikCreateInputSerializer

    def post(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        excel_bytes, period_label = generate_laporan_statistik_kelompok_excel(
            bulan=data['bulan'],
            tahun=data['tahun'],
        )

        with transaction.atomic():
            laporan = serializer.save()
            excel_filename = f"laporan_statistik_kelompok{period_label}.xlsx"
            laporan.file_excel.save(excel_filename, ContentFile(excel_bytes), save=True)

        output_serializer = LaporanStatistikSerializer(laporan, context={'request': request})
        return Response(output_serializer.data, status=status.HTTP_201_CREATED)


class LaporanStatistikKelompokExcelView(generics.GenericAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]

    def get(self, request, *args, **kwargs):
        bulan = request.query_params.get('bulan')
        tahun = request.query_params.get('tahun')
        excel_bytes, period_label = generate_laporan_statistik_kelompok_excel(
            bulan=bulan,
            tahun=tahun,
        )

        response = HttpResponse(
            excel_bytes,
            content_type=XLSX_CONTENT_TYPE,
        )
        response['Content-Disposition'] = (
            f'attachment; filename="laporan_statistik_kelompok{period_label}.xlsx"'
        )
        return response


class LaporanSTDBExcelView(generics.GenericAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]

    @staticmethod
    def _raw_or_empty(obj, field_name, fallback=''):
        if obj is None:
            return fallback

        value = getattr(obj, field_name, None)

        if value in (None, ''):
            return fallback
        return value

    @staticmethod
    def _komoditas_value(komoditas):
        if komoditas in (None, ''):
            return ''
        return json.dumps([str(komoditas)], separators=(',', ':'))

    @staticmethod
    def _date_or_empty(date_value):
        if not date_value:
            return ''
        return date_value.strftime('%Y-%m-%d')

    @staticmethod
    def _luas_to_m2(luas_ha):
        if luas_ha in (None, ''):
            return ''

        m2 = Decimal(str(luas_ha)) * Decimal('10000')
        return int(m2) if m2 == m2.to_integral_value() else float(m2)

    @staticmethod
    def _geom_to_text(geom):
        if not geom:
            return ''

        coords = getattr(geom, 'coords', None)
        if not coords:
            return ''

        # PolygonField stores rings; use exterior ring to match requested format.
        ring = coords[0] if coords and coords[0] else []
        points = [[float(point[0]), float(point[1])] for point in ring]
        return json.dumps(points, separators=(',', ':'))

    @staticmethod
    def _build_kebun_queryset(status_param):
        kebun_terbit_q = Q(nomor_stdb__isnull=False) & ~Q(nomor_stdb='')
        kebun_belum_q = Q(nomor_stdb__isnull=True) | Q(nomor_stdb='')

        kebun_qs = Kebun.objects.select_related('petani', 'petani__desa', 'desa').order_by(
            'petani__nama_kelompok',
            'petani__nama',
            'id_kebun',
        )
        return kebun_qs.filter(kebun_terbit_q if status_param == 'terbit' else kebun_belum_q)

    def _build_row_values(self, kebun):
        petani = kebun.petani
        return [
            petani.nama or '',
            petani.tempat or '',
            self._date_or_empty(petani.tanggal_lahir),
            petani.no_ktp or '',
            petani.alamat or '',
            petani.desa_id or '',
            self._raw_or_empty(petani, 'jns_kelamin'),
            self._raw_or_empty(petani, 'pendidikan_terakhir'),
            self._raw_or_empty(kebun, 'jenis_legalitas'),
            kebun.nomor_legalitas or '',
            self._komoditas_value(kebun.komoditas),
            petani.nama or '',
            self._luas_to_m2(kebun.luas),
            kebun.total_prod_per_tahun if kebun.total_prod_per_tahun is not None else '',
            kebun.waktu_tanam.year if kebun.waktu_tanam else '',
            kebun.tahun_peremajaan if kebun.tahun_peremajaan is not None else '',
            kebun.jumlah_pohon if kebun.jumlah_pohon is not None else '',
            self._raw_or_empty(kebun, 'pola_tanam'),
            self._raw_or_empty(kebun, 'jenis_lahan'),
            self._raw_or_empty(kebun, 'asal_benih'),
            self._raw_or_empty(kebun, 'jenis_pupuk'),
            kebun.mitra_penjualan or '',
            kebun.desa_id or '',
            '',
            kebun.nomor_stdb or '',
            self._geom_to_text(kebun.geom),
        ]

    def get(self, request, *args, **kwargs):
        status_param = request.query_params.get('status')
        if status_param not in ('terbit', 'belum'):
            return Response(
                {'detail': "Parameter 'status' wajib diisi dengan nilai: terbit | belum."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        kebun_qs = self._build_kebun_queryset(status_param)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Laporan STDB Petani'

        headers = [
            'nama',
            'tempat_lahir',
            'tanggal_lahir',
            'NIK',
            'alamat',
            'kw_domisili',
            'jenis_kelamin',
            'ijazah_terakhir',
            'leg_status_kepemilikan',
            'leg_nomor',
            'kom_jenis_komoditi_kode',
            'nama_sesuai_sertifikat',
            'luas_lahan_sertifikat',
            'prod_total_prod_setahun',
            'kom_tahun_tanam',
            'kom_tahun_tanam_peremajaan',
            'kom_jml_pohon',
            'kom_pola_tanam',
            'kon_jenis_lahan',
            'prod_asal_benih',
            'prod_jenis_pupuk',
            'mitra_pengolahan',
            'kw_kelurahan',
            'tanggal_terbit',
            'no_stdb',
            'map_geoshape',
        ]

        header_font = Font(bold=True)
        for col, title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font

        for row_idx, kebun in enumerate(kebun_qs, start=2):
            for col, value in enumerate(self._build_row_values(kebun), 1):
                ws.cell(row=row_idx, column=col, value=value)

        ws.freeze_panes = 'A2'

        response = HttpResponse(
            content_type=XLSX_CONTENT_TYPE,
        )
        response['Content-Disposition'] = (
            f'attachment; filename="laporan_stdb_petani_{status_param}.xlsx"'
        )

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        response.write(output.getvalue())
        return response


class LaporanPetaniView(generics.GenericAPIView):
    authentication_classes = [JWTAuthentication]
    permission_classes = [IsAuthenticated, role_required(1, 3, 4, 5)]
    serializer_class = LaporanPetaniSerializer

    def get(self, request, id_petani, *args, **kwargs):
        petani = get_object_or_404(Petani, pk=id_petani)
        kebun_qs = Kebun.objects.filter(petani_id=id_petani).select_related('petani')
        pekerja_qs = Pekerja.objects.filter(petani_id=id_petani).select_related('petani')
        diklat_qs = Diklat.objects.filter(petani_id=id_petani).select_related('petani')
        # kelompok_penyetor_qs = KelompokPenyetor.objects.filter(
        #     anggota_petani__id=id_petani,
        #     angkutan__isnull=False,
        # ).select_related('angkutan').distinct()
        # penjualan_qs = Angkutan.objects.filter(
        #     id__in=kelompok_penyetor_qs.values_list('angkutan_id', flat=True)
        # ).select_related('pabrik')

        serializer = self.get_serializer(
            {
                'data_petani': petani,
                'data_kebun': kebun_qs,
                'data_produksi': kebun_qs,
                'data_pestisida': kebun_qs,
                'data_pupuk': kebun_qs,
                'data_lb3': kebun_qs,
                'data_pekerja': pekerja_qs,
                'data_diklat': diklat_qs,
                # 'data_penjualan': penjualan_qs,
            },
            context={'request': request},
        )
        return Response(serializer.data, status=status.HTTP_200_OK)
