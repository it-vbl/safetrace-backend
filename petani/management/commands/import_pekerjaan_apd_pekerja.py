import re
from pathlib import Path

from openpyxl import load_workbook

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from petani.models import Pekerja
from utils.choices import JenisAPD, JenisPekerjaan


PEKERJAAN_COLUMNS = (
    (JenisPekerjaan.PANEN, ('panen',)),
    (JenisPekerjaan.PUPUK, ('pupuk',)),
    (JenisPekerjaan.SEMPROT, ('semprot',)),
    (JenisPekerjaan.TEBAS, ('tebas',)),
    (JenisPekerjaan.SUPIR, ('supir',)),
    (JenisPekerjaan.LANGSIR, ('langsir',)),
)

APD_COLUMNS = (
    (JenisAPD.APRON, ('apron',)),
    (JenisAPD.KACA_MATA, ('kaca mata',)),
    (JenisAPD.SEPATU, ('sepatu',)),
    (JenisAPD.HELEM, ('helem',)),
    (JenisAPD.MASKER, ('masker',)),
    (JenisAPD.SARUNG_TANGAN, ('sarung tangan',)),
    (JenisAPD.SARUNG_EGREK_DODOS, ('sarung egrek dodos', 'sarung egrek/dodos')),
)


class Command(BaseCommand):
    help = 'Import data jenis pekerjaan dan jenis APD pekerja dari file Excel'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Path ke file Excel')
        parser.add_argument(
            '--header-row-1',
            type=int,
            default=1,
            help='Nomor baris header pertama (default: 1)',
        )
        parser.add_argument(
            '--header-row-2',
            type=int,
            default=2,
            help='Nomor baris header kedua (default: 2)',
        )
        parser.add_argument(
            '--data-start-row',
            type=int,
            default=0,
            help='Nomor baris awal data. Jika 0, otomatis header_row_2 + 1',
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Jalankan tanpa menyimpan data (testing mode)',
        )

    def normalize_text(self, value):
        if value is None:
            return ''

        if isinstance(value, bool):
            return '1' if value else '0'

        if isinstance(value, int):
            return str(value)

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return str(value)

        return str(value).strip()

    def normalize_header(self, value):
        text = self.normalize_text(value).lower()
        text = re.sub(r'[^0-9a-z]+', ' ', text)
        return re.sub(r'\s+', ' ', text).strip()

    def is_one(self, value):
        if value is None:
            return False

        if isinstance(value, bool):
            return value

        if isinstance(value, (int, float)):
            return float(value) == 1

        normalized = self.normalize_text(value).strip().lower()
        return normalized in {'1', '1.0', 'true', 'yes', 'y'}

    def normalize_identifier(self, value):
        text = self.normalize_text(value).strip()
        if not text:
            return ''

        if re.fullmatch(r'\d+\.0', text):
            return text[:-2]

        return text

    def build_header_map(self, worksheet, header_row_1, header_row_2):
        header_map = {}

        for column_index in range(1, worksheet.max_column + 1):
            top_header = self.normalize_header(worksheet.cell(row=header_row_1, column=column_index).value)
            second_header = self.normalize_header(worksheet.cell(row=header_row_2, column=column_index).value)

            header = second_header or top_header
            if header:
                header_map[header] = column_index

        return header_map

    def detect_header_rows(self, worksheet, max_scan_rows=40):
        last_row_to_scan = min(worksheet.max_row - 1, max_scan_rows)
        if last_row_to_scan < 1:
            return None, None, None, None, None

        for row_1 in range(1, last_row_to_scan + 1):
            row_2 = row_1 + 1
            if row_2 > worksheet.max_row:
                break

            header_map = self.build_header_map(worksheet, row_1, row_2)
            try:
                no_ktp_col, pekerjaan_columns, apd_columns = self.resolve_required_columns(header_map)
                return row_1, row_2, no_ktp_col, pekerjaan_columns, apd_columns
            except CommandError:
                continue

        return None, None, None, None, None

    def resolve_required_columns(self, header_map):
        no_ktp_col = self.resolve_column(header_map, ('no nik', 'no ktp', 'nik'), 'No NIK')

        pekerjaan_columns = []
        for choice_value, aliases in PEKERJAAN_COLUMNS:
            column_index = self.resolve_column(header_map, aliases, choice_value.label)
            pekerjaan_columns.append((choice_value, column_index))

        apd_columns = []
        for choice_value, aliases in APD_COLUMNS:
            column_index = self.resolve_column(header_map, aliases, choice_value.label)
            apd_columns.append((choice_value, column_index))

        return no_ktp_col, pekerjaan_columns, apd_columns

    def validate_row_options(self, header_row_1, header_row_2, data_start_row):
        if header_row_1 < 1 or header_row_2 < 1:
            raise CommandError('header-row-1 dan header-row-2 harus >= 1')

        if header_row_2 <= header_row_1:
            raise CommandError('header-row-2 harus lebih besar dari header-row-1')

        if data_start_row and data_start_row <= header_row_2:
            raise CommandError('data-start-row harus lebih besar dari header-row-2')

    def resolve_headers(self, worksheet, header_row_1, header_row_2):
        self.header_map = self.build_header_map(worksheet, header_row_1, header_row_2)

        try:
            no_ktp_column, pekerjaan_columns, apd_columns = self.resolve_required_columns(self.header_map)
            return header_row_1, header_row_2, no_ktp_column, pekerjaan_columns, apd_columns
        except CommandError:
            detected_row_1, detected_row_2, detected_no_ktp, pekerjaan_columns, apd_columns = self.detect_header_rows(
                worksheet
            )
            if detected_row_1 is None:
                raise

            self.stdout.write(
                self.style.WARNING(
                    f'Header otomatis ditemukan di baris {detected_row_1} dan {detected_row_2}'
                )
            )
            self.header_map = self.build_header_map(worksheet, detected_row_1, detected_row_2)
            return detected_row_1, detected_row_2, detected_no_ktp, pekerjaan_columns, apd_columns

    def get_effective_data_start_row(self, data_start_row, header_row_2):
        return data_start_row if data_start_row > 0 else header_row_2 + 1

    def resolve_column(self, header_map, aliases, field_label):
        for alias in aliases:
            normalized_alias = self.normalize_header(alias)
            column_index = header_map.get(normalized_alias)
            if column_index:
                return column_index

        raise CommandError(f'Kolom "{field_label}" tidak ditemukan pada header Excel')

    def get_row_value(self, row_values, column_index):
        if column_index <= 0:
            return None

        row_offset = column_index - 1
        if row_offset >= len(row_values):
            return None

        return row_values[row_offset]

    def collect_selected_values(self, row_values, resolved_columns):
        selected_values = []

        for choice_value, column_index in resolved_columns:
            if self.is_one(self.get_row_value(row_values, column_index)):
                selected_values.append(choice_value)

        return selected_values

    def process_row(self, row_num, row_values, dry_run):
        no_ktp = self.normalize_identifier(self.get_row_value(row_values, self.no_ktp_column))

        if not no_ktp:
            self.skip_count += 1
            self.stdout.write(self.style.WARNING(f"⊘ Baris {row_num}: No NIK kosong, dilewati"))
            return

        pekerja_qs = Pekerja.objects.filter(no_ktp=no_ktp).order_by('id')
        pekerja_count = pekerja_qs.count()

        if pekerja_count == 0:
            self.skip_count += 1
            self.stdout.write(
                self.style.WARNING(
                    f"⊘ Baris {row_num}: Pekerja dengan No NIK '{no_ktp}' tidak ditemukan, dilewati"
                )
            )
            return

        if pekerja_count > 1:
            self.error_count += 1
            self.stdout.write(
                self.style.ERROR(
                    f"✗ Baris {row_num}: Ditemukan {pekerja_count} pekerja dengan No NIK '{no_ktp}'"
                )
            )
            return

        pekerja = pekerja_qs.first()
        jenis_pekerjaan = self.collect_selected_values(row_values, self.pekerjaan_column_indexes)
        jenis_apd = self.collect_selected_values(row_values, self.apd_column_indexes)

        changed = False
        if pekerja.jenis_pekerjaan != jenis_pekerjaan:
            pekerja.jenis_pekerjaan = jenis_pekerjaan
            changed = True

        if pekerja.jenis_apd != jenis_apd:
            pekerja.jenis_apd = jenis_apd
            changed = True

        if not dry_run and changed:
            pekerja.save(update_fields=['jenis_pekerjaan', 'jenis_apd', 'updated_at'])

        self.success_count += 1

        jenis_pekerjaan_label = ', '.join(label for _, label in JenisPekerjaan.choices if _ in jenis_pekerjaan) or '-'
        jenis_apd_label = ', '.join(label for _, label in JenisAPD.choices if _ in jenis_apd) or '-'
        if dry_run:
            action = 'akan diupdate'
        elif changed:
            action = 'diupdate'
        else:
            action = 'tidak berubah'

        self.stdout.write(
            self.style.SUCCESS(
                f"✓ Baris {row_num}: {pekerja.nama} ({no_ktp}) {action} | "
                f"Jenis Pekerjaan: {jenis_pekerjaan_label} | Jenis APD: {jenis_apd_label}"
            )
        )

    def print_summary(self, dry_run):
        self.stdout.write('\n' + '=' * 70)
        self.stdout.write(self.style.SUCCESS(f'Berhasil: {self.success_count} record'))
        self.stdout.write(self.style.WARNING(f'Dilewati: {self.skip_count} record'))
        self.stdout.write(self.style.ERROR(f'Gagal: {self.error_count} record'))

        if dry_run:
            self.stdout.write(self.style.WARNING('\n⚠ DRY RUN MODE - Tidak ada data yang disimpan'))
        else:
            self.stdout.write(
                self.style.SUCCESS(
                    f'\n✓ Import selesai! Total {self.success_count} data pekerja berhasil diproses'
                )
            )

        if self.errors:
            self.stdout.write('\n' + self.style.ERROR('Daftar error:'))
            for error in self.errors[:20]:
                self.stdout.write(f'  {error}')
            if len(self.errors) > 20:
                self.stdout.write(f'  ... dan {len(self.errors) - 20} error lainnya')

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        dry_run = options['dry_run']
        header_row_1 = options['header_row_1']
        header_row_2 = options['header_row_2']
        data_start_row = options['data_start_row']

        self.success_count = 0
        self.skip_count = 0
        self.error_count = 0
        self.errors = []
        self.header_map = {}
        self.no_ktp_column = None
        self.pekerjaan_column_indexes = []
        self.apd_column_indexes = []
        self.pekerjaan_columns = PEKERJAAN_COLUMNS
        self.apd_columns = APD_COLUMNS

        if dry_run:
            self.stdout.write(self.style.WARNING('Running in DRY RUN mode - data will not be saved'))

        excel_path = Path(excel_file)
        if not excel_path.exists():
            raise CommandError(f'File tidak ditemukan: {excel_file}')

        self.validate_row_options(header_row_1, header_row_2, data_start_row)

        try:
            workbook = load_workbook(excel_path, data_only=True)
            worksheet = workbook.worksheets[0]

            if worksheet.max_row < header_row_2 + 1:
                raise CommandError('Sheet pertama tidak memiliki data yang cukup untuk diproses')

            (
                header_row_1,
                header_row_2,
                self.no_ktp_column,
                self.pekerjaan_column_indexes,
                self.apd_column_indexes,
            ) = self.resolve_headers(
                worksheet,
                header_row_1,
                header_row_2,
            )

            data_start_row = self.get_effective_data_start_row(data_start_row, header_row_2)

            if data_start_row <= header_row_2:
                raise CommandError('data-start-row harus lebih besar dari header-row-2')

            with transaction.atomic():
                for row_num, row_values in enumerate(
                    worksheet.iter_rows(min_row=data_start_row, values_only=True),
                    start=data_start_row,
                ):
                    try:
                        if not any(value not in (None, '') for value in row_values):
                            self.skip_count += 1
                            continue

                        self.process_row(row_num, row_values, dry_run)
                    except Exception as exc:
                        self.error_count += 1
                        error_msg = f"✗ Baris {row_num}: {str(exc)}"
                        self.errors.append(error_msg)
                        self.stdout.write(self.style.ERROR(error_msg))

                if dry_run:
                    transaction.set_rollback(True)

            self.print_summary(dry_run)

        except CommandError:
            raise
        except Exception as exc:
            raise CommandError(f'Error membaca file Excel: {str(exc)}')