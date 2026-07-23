from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from gap.models import Produksi
from petani.models import Petani


class Command(BaseCommand):
    help = 'Import data produksi GAP dari file Excel (sheet kedua) ke model Produksi'

    HEADER_ROW_1 = 10
    HEADER_ROW_2 = 11
    DATA_START_ROW = 12

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Path ke file Excel')
        parser.add_argument('tahun', type=int, help='Tahun data produksi')
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Jalankan tanpa menyimpan data (testing mode)',
        )

    def normalize_text(self, value):
        if value is None:
            return ''

        if isinstance(value, bool):
            return '1' if value else '0'

        if isinstance(value, int):
            return str(value)

        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return str(value)

        return str(value).strip()

    def normalize_header(self, value):
        text = self.normalize_text(value).lower()
        text = re.sub(r'[^0-9a-z]+', ' ', text)
        return re.sub(r'\s+', ' ', text).strip()

    def normalize_name(self, value):
        text = self.normalize_text(value)
        return re.sub(r'\s+', ' ', text).strip()

    def parse_decimal(self, value):
        if value in (None, ''):
            return Decimal('0')

        if isinstance(value, Decimal):
            return value

        if isinstance(value, (int, float)):
            return Decimal(str(value))

        text = self.normalize_text(value)
        if not text:
            return Decimal('0')

        text = text.replace(' ', '')
        text = text.replace('.', '') if ',' in text and '.' in text else text
        text = text.replace(',', '.')

        try:
            return Decimal(text)
        except (InvalidOperation, ValueError):
            raise ValueError(f'Nilai numerik tidak valid: {value}')

    def build_header_map(self, worksheet):
        header_map = {}

        for column_index in range(1, worksheet.max_column + 1):
            top_header = self.normalize_header(
                worksheet.cell(row=self.HEADER_ROW_1, column=column_index).value
            )
            second_header = self.normalize_header(
                worksheet.cell(row=self.HEADER_ROW_2, column=column_index).value
            )

            header = second_header or top_header
            if header:
                header_map[header] = column_index

        return header_map

    def resolve_column(self, header_map, aliases, field_label):
        for alias in aliases:
            normalized_alias = self.normalize_header(alias)
            column_index = header_map.get(normalized_alias)
            if column_index:
                return column_index

        raise CommandError(f'Kolom "{field_label}" tidak ditemukan pada header Excel')

    def get_row_value(self, row_values, column_index):
        if column_index <= 0:
            return None

        row_offset = column_index - 1
        if row_offset >= len(row_values):
            return None

        return row_values[row_offset]

    def resolve_columns(self, header_map):
        return {
            'nama_petani': self.resolve_column(header_map, ('nama petani',), 'NAMA PETANI'),
            'luas_kebun': self.resolve_column(header_map, ('luas kebun',), 'LUAS KEBUN'),
            'tahun_tanam': self.resolve_column(header_map, ('tahun tanam',), 'TAHUN TANAM'),
            'januari': self.resolve_column(header_map, ('januari',), 'Januari'),
            'februari': self.resolve_column(header_map, ('februari',), 'Februari'),
            'maret': self.resolve_column(header_map, ('maret',), 'Maret'),
            'april': self.resolve_column(header_map, ('april',), 'April'),
            'mei': self.resolve_column(header_map, ('mei',), 'Mei'),
            'juni': self.resolve_column(header_map, ('juni',), 'Juni'),
            'juli': self.resolve_column(header_map, ('juli',), 'Juli'),
            'agustus': self.resolve_column(header_map, ('agustus',), 'Agustus'),
            'september': self.resolve_column(header_map, ('september',), 'September'),
            'oktober': self.resolve_column(header_map, ('oktober',), 'Oktober'),
            'november': self.resolve_column(header_map, ('november',), 'November'),
            'desember': self.resolve_column(header_map, ('desember',), 'Desember'),
        }

    def find_petani(self, nama_petani):
        return Petani.objects.filter(nama__iexact=nama_petani).order_by('id').first()

    def process_row(self, row_num, row_values, tahun, dry_run):
        nama_petani = self.normalize_name(
            self.get_row_value(row_values, self.columns['nama_petani'])
        )

        if not nama_petani:
            self.skip_count += 1
            return

        if nama_petani.lower() == 'total':
            self.skip_count += 1
            return

        if self.previous_name and nama_petani.lower() == self.previous_name.lower():
            petani = self.previous_petani
            kebun_sequence = self.previous_kebun_sequence + 1
        else:
            petani = self.find_petani(nama_petani)
            kebun_sequence = 1

        self.previous_name = nama_petani
        self.previous_petani = petani
        self.previous_kebun_sequence = kebun_sequence

        if not petani:
            self.skip_count += 1
            self.stdout.write(
                self.style.WARNING(
                    f"⊘ Baris {row_num}: Petani '{nama_petani}' tidak ditemukan, dilewati"
                )
            )
            return

        kebun = petani.kebun_set.order_by('id')[kebun_sequence - 1:kebun_sequence].first()
        if not kebun:
            self.skip_count += 1
            self.stdout.write(
                self.style.WARNING(
                    f"⊘ Baris {row_num}: Kebun ke-{kebun_sequence} untuk petani '{nama_petani}' tidak ditemukan, dilewati"
                )
            )
            return

        produksi_payload = {
            'kebun': kebun,
            'tahun': tahun,
            'januari': self.parse_decimal(self.get_row_value(row_values, self.columns['januari'])),
            'februari': self.parse_decimal(self.get_row_value(row_values, self.columns['februari'])),
            'maret': self.parse_decimal(self.get_row_value(row_values, self.columns['maret'])),
            'april': self.parse_decimal(self.get_row_value(row_values, self.columns['april'])),
            'mei': self.parse_decimal(self.get_row_value(row_values, self.columns['mei'])),
            'juni': self.parse_decimal(self.get_row_value(row_values, self.columns['juni'])),
            'juli': self.parse_decimal(self.get_row_value(row_values, self.columns['juli'])),
            'agustus': self.parse_decimal(self.get_row_value(row_values, self.columns['agustus'])),
            'september': self.parse_decimal(self.get_row_value(row_values, self.columns['september'])),
            'oktober': self.parse_decimal(self.get_row_value(row_values, self.columns['oktober'])),
            'november': self.parse_decimal(self.get_row_value(row_values, self.columns['november'])),
            'desember': self.parse_decimal(self.get_row_value(row_values, self.columns['desember'])),
        }

        if not dry_run:
            Produksi.objects.create(**produksi_payload)

        self.success_count += 1
        status = 'akan diinsert' if dry_run else 'diinsert'
        self.stdout.write(
            self.style.SUCCESS(
                f"✓ Baris {row_num}: {nama_petani} | kebun ke-{kebun_sequence} | {status}"
            )
        )

    def print_summary(self, dry_run):
        self.stdout.write('\n' + '=' * 70)
        self.stdout.write(self.style.SUCCESS(f'Berhasil: {self.success_count} record'))
        self.stdout.write(self.style.WARNING(f'Dilewati: {self.skip_count} record'))
        self.stdout.write(self.style.ERROR(f'Gagal: {self.error_count} record'))

        if dry_run:
            self.stdout.write(self.style.WARNING('\n⚠ DRY RUN MODE - Tidak ada data yang disimpan'))
        else:
            self.stdout.write(self.style.SUCCESS('\n✓ Import selesai! Data produksi berhasil diproses'))

        if self.errors:
            self.stdout.write('\n' + self.style.ERROR('Daftar error:'))
            for error in self.errors[:20]:
                self.stdout.write(f'  {error}')
            if len(self.errors) > 20:
                self.stdout.write(f'  ... dan {len(self.errors) - 20} error lainnya')

    def handle(self, *args, **options):
        excel_file = options['excel_file']
        tahun = options['tahun']
        dry_run = options['dry_run']

        self.success_count = 0
        self.skip_count = 0
        self.error_count = 0
        self.errors = []
        self.previous_name = None
        self.previous_petani = None
        self.previous_kebun_sequence = 0
        self.columns = {}

        if dry_run:
            self.stdout.write(self.style.WARNING('Running in DRY RUN mode - data will not be saved'))

        excel_path = Path(excel_file)
        if not excel_path.exists():
            raise CommandError(f'File tidak ditemukan: {excel_file}')

        try:
            workbook = load_workbook(excel_path, data_only=True)
            if len(workbook.worksheets) < 2:
                raise CommandError('File Excel harus memiliki minimal 2 sheet')

            worksheet = workbook.worksheets[1]

            if worksheet.max_row < self.DATA_START_ROW:
                raise CommandError('Sheet kedua tidak memiliki data yang cukup untuk diproses')

            header_map = self.build_header_map(worksheet)
            self.columns = self.resolve_columns(header_map)

            with transaction.atomic():
                for row_num, row_values in enumerate(
                    worksheet.iter_rows(min_row=self.DATA_START_ROW, values_only=True),
                    start=self.DATA_START_ROW,
                ):
                    try:
                        if not any(value not in (None, '') for value in row_values):
                            self.skip_count += 1
                            continue

                        self.process_row(row_num, row_values, tahun, dry_run)
                    except Exception as exc:
                        self.error_count += 1
                        error_msg = f"✗ Baris {row_num}: {str(exc)}"
                        self.errors.append(error_msg)
                        self.stdout.write(self.style.ERROR(error_msg))

                if dry_run:
                    transaction.set_rollback(True)

            self.print_summary(dry_run)

        except CommandError:
            raise
        except Exception as exc:
            raise CommandError(f'Error membaca file Excel: {str(exc)}')
